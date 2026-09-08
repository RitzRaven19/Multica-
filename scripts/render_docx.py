#!/usr/bin/env python3
"""Render BRD.docx, FRD.docx, UAT.docx from the approved Group 2 pipeline output.

This is a deterministic renderer, not another agent pass: every Word
document is built directly from shared_outputs/*.json (and BRD.md for
the BRD, since brd-agent already writes that as prose). That keeps the
.docx files a rendering of the approved data, never a second,
independently-generated account of it that could drift out of sync --
the failure mode the earlier markdown-only pipeline hit with a stale
FRD.

Requires: pip install python-docx
"""
import csv
import json
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent.parent / "shared_outputs"


def load_json(name):
    with open(OUT / name, encoding="utf-8-sig") as f:
        return json.load(f)


def load_csv(name):
    with open(OUT / name, encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if val is None else str(val)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    return table


def status_line(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


# ---------------------------------------------------------------- BRD ----

_INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_inline(paragraph, text):
    pos = 0
    for m in _INLINE_BOLD.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def render_brd():
    md_path = OUT / "BRD.md"
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    doc.add_heading("Business Requirements Document", level=0)

    i = 0
    table_buf = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|"):
            table_buf.append(stripped)
            i += 1
            continue
        elif table_buf:
            rows = [
                [c.strip() for c in r.strip("|").split("|")]
                for r in table_buf
                if not re.match(r"^\|?\s*-+\s*\|", r)
            ]
            if rows:
                add_table(doc, rows[0], rows[1:])
            table_buf = []

        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, stripped[2:])
        elif stripped:
            p = doc.add_paragraph()
            _add_inline(p, stripped)
        i += 1

    if table_buf:
        rows = [
            [c.strip() for c in r.strip("|").split("|")]
            for r in table_buf
            if not re.match(r"^\|?\s*-+\s*\|", r)
        ]
        if rows:
            add_table(doc, rows[0], rows[1:])

    doc.save(OUT / "BRD.docx")
    print("wrote BRD.docx")


# ---------------------------------------------------------------- FRD ----

def render_frd():
    frd = load_json("frd.json")
    data_spec = load_json("data_spec.json")
    errcat = load_json("error_catalogue.json")

    doc = Document()
    doc.add_heading("Functional / Software Requirements Specification", level=0)
    status_line(doc, "Status: DRAFT-derived from approved shared_outputs/frd.json, data_spec.json, error_catalogue.json")

    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        f"This FRD covers {len(frd)} requirements (FR-001 through FR-{len(frd):03d}), each traced to a "
        "BRD functional or non-functional requirement id and one or more business rules. "
        "Rendered directly from the approved pipeline data -- see traceability_matrix.csv for the full chain."
    )

    doc.add_heading("2. Requirements", level=1)
    for f in frd:
        doc.add_heading(f"{f['requirement_id']} — {f['title']}  (traces {f['official_id']})", level=2)
        p = doc.add_paragraph()
        p.add_run("Statement: ").bold = True
        p.add_run(f["statement"])

        meta = doc.add_paragraph()
        meta.add_run("Epic: ").bold = True
        meta.add_run(f"{f['epic']}    ")
        meta.add_run("Priority: ").bold = True
        meta.add_run(f"{f['priority']} ({f['priority_code']})    ")
        meta.add_run("Roles: ").bold = True
        meta.add_run(", ".join(f.get("roles", [])))

        if f.get("preconditions"):
            doc.add_paragraph("Preconditions:", style="Intense Quote")
            for pc in f["preconditions"]:
                doc.add_paragraph(pc, style="List Bullet")
        if f.get("postconditions"):
            doc.add_paragraph("Postconditions:", style="Intense Quote")
            for pc in f["postconditions"]:
                doc.add_paragraph(pc, style="List Bullet")
        if f.get("error_ids"):
            doc.add_paragraph("Errors: " + ", ".join(f["error_ids"]))
        if f.get("open_question_ids"):
            doc.add_paragraph("Blocked by open question(s): " + ", ".join(f["open_question_ids"]))
        doc.add_paragraph(f"Source: {', '.join(f.get('source_ids', []))}")

    doc.add_heading("3. Data Specification", level=1)
    doc.add_heading("3.1 Entity Fields", level=2)
    add_table(
        doc,
        ["Entity", "Field", "Type", "Format", "Length", "Example"],
        [[fl["entity"], fl["name"], fl["type"], fl.get("format", ""), fl.get("length", ""), fl.get("example", "")]
         for fl in data_spec["fields"]],
    )

    doc.add_heading("3.2 State Transitions", level=2)
    for entity, states in data_spec["state_transitions"].items():
        p = doc.add_paragraph()
        p.add_run(f"{entity}: ").bold = True
        p.add_run(" → ".join(states) if isinstance(states, list) else str(states))

    if data_spec.get("limit_placeholders"):
        doc.add_heading("3.3 Numeric Limit Placeholders", level=2)
        add_table(
            doc,
            ["Name", "Blocked by", "Below", "At", "Above", "Error"],
            [[lp["name"], lp.get("open_question_id", ""),
              lp["boundary_values"]["below"], lp["boundary_values"]["at"], lp["boundary_values"]["above"],
              lp.get("error_id", "")]
             for lp in data_spec["limit_placeholders"]],
        )

    doc.add_heading("4. Error Catalogue", level=1)
    add_table(
        doc,
        ["Error ID", "Name", "HTTP", "User Message", "Notes"],
        [[e["error_id"], e["name"], e["http_status"], e["user_message"], e.get("notes", "")] for e in errcat],
    )

    doc.save(OUT / "FRD.docx")
    print("wrote FRD.docx")


# ---------------------------------------------------------------- UAT ----

def render_uat():
    ac = load_json("acceptance_criteria.json")
    uat = load_json("uat_cases.json")
    cov = load_json("coverage_report.json")
    matrix_rows = load_csv("traceability_matrix.csv")

    doc = Document()
    doc.add_heading("User Acceptance Test Plan", level=0)
    status_line(doc, "Status: DRAFT-derived from approved shared_outputs/acceptance_criteria.json, uat_cases.json, traceability_matrix.csv, coverage_report.json")

    doc.add_heading("1. Coverage Summary", level=1)
    p = doc.add_paragraph()
    p.add_run("Overall status: ").bold = True
    run = p.add_run(cov["status"])
    run.bold = True
    add_table(
        doc,
        ["Metric", "Count"],
        [[k, v] for k, v in cov["counts"].items()],
    )
    add_table(
        doc,
        ["Acceptance Rate", "Value"],
        [[k, v] for k, v in cov["acceptance_rate"].items()],
    )
    if cov.get("issues"):
        doc.add_heading("Open Issues", level=2)
        for issue in cov["issues"]:
            doc.add_paragraph(issue, style="List Bullet")

    doc.add_heading("2. Acceptance Criteria", level=1)
    add_table(
        doc,
        ["AC ID", "Requirement", "Type", "Given", "When", "Then"],
        [[a["ac_id"], a["requirement_id"], a["type"], a["given"], a["when"], a["then"]] for a in ac],
    )

    doc.add_heading("3. Test Cases", level=1)
    for u in uat:
        doc.add_heading(f"{u['uat_id']} — {u['title']}", level=2)
        p = doc.add_paragraph()
        p.add_run("Role: ").bold = True
        p.add_run(f"{u['role']}    ")
        p.add_run("Requirement(s): ").bold = True
        p.add_run(", ".join(u.get("requirement_ids", [])))
        if u.get("preconditions"):
            doc.add_paragraph("Preconditions: " + "; ".join(u["preconditions"]))
        for step in u.get("steps", []):
            doc.add_paragraph(f"{step['step']}. {step['action']} → {step['expected']}", style="List Number")
        doc.add_paragraph("Result: (blank -- filled in during execution)")

    doc.add_heading("4. Traceability Matrix", level=1)
    if matrix_rows:
        add_table(doc, matrix_rows[0], matrix_rows[1:])

    doc.save(OUT / "UAT.docx")
    print("wrote UAT.docx")


if __name__ == "__main__":
    render_brd()
    render_frd()
    render_uat()
